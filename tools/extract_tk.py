#!/usr/bin/env python3
"""Zet het EASA "TK Syllabus Comparison Document" (xlsx) om naar JSON: per leerdoel de syllabusreferentie,
de tekst, of het BK (basiskennis, geen aparte examenvraag) is, voor welke brevetten het geldt en de bron.

    python3 tools/extract_tk.py <tk-syllabus-comparison-doc-v6.xlsx>

Let op: dit is de syllabus voor ATPL/CPL/IR (Part-FCL, AMC1 FCL.310 …). De leerdoelen voor LAPL/PPL zijn een apart
document, maar veel onderwerpen (en dus de bronnen) zijn dezelfde.
"""
import json
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "sources", "ecqb")


def main(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    info = {}
    for row in wb["Reader Instructions"].iter_rows(values_only=True):
        if row[1] == "Version:": info["versie"] = row[2]
        if row[1] == "Date:": info["datum"] = str(row[2])[:10]
        if row[1] == "For ECQB release:": info["ecqb"] = row[2]
    bronnen = []
    for row in list(wb["Source-information"].iter_rows(values_only=True))[4:]:
        if row[0]:
            bronnen.append({"bron": str(row[0]).strip(), "versie": str(row[1] or "").strip(), "link": str(row[2] or "").strip()})
    vakken = {}
    for ws in wb.worksheets[2:]:
        rows = list(ws.iter_rows(values_only=True))
        head = [str(h or "").strip() for h in rows[0]]
        col = {h: i for i, h in enumerate(head)}
        licences = [h for h in head[12:] if h and not h.startswith(("Source", "Entry", "BIR BK"))]
        src_col = next(i for i, h in enumerate(head) if h.startswith("Source"))
        ver_col = next((i for i, h in enumerate(head) if h.startswith("Entry")), None)
        items = []
        for r in rows[1:]:
            ref = r[col["2020 syllabus reference"]]
            if not ref:
                continue
            ref = str(ref).strip()
            items.append({
                "ref": ref,
                "niveau": ref.count(".") + 1,
                "tekst": " ".join(str(r[col["2020 syllabus text"]] or "").split()),
                "bk": bool(r[col["BK"]]),
                "brevetten": [l for l in licences if r[col[l]]],
                "bron": " ".join(str(r[src_col] or "").split()),
                "tk_versie": r[ver_col] if ver_col is not None else None,
            })
        vakken[ws.title] = items
    os.makedirs(OUT, exist_ok=True)
    data = {
        "titel": "EASA TK Syllabus Comparison Document",
        **info,
        "bron": "EASA – https://www.easa.europa.eu/en/domains/aircrew-and-medical/european-central-question-bank-ecqb",
        "opmerking": "Syllabus en leerdoelen voor ATPL/CPL/IR (ED Decision 2020/018/R) met per leerdoel de bron voor ECQB 2026. "
                     "Niet de LAPL/PPL-leerdoelen, maar veel onderwerpen en bronnen zijn gelijk.",
        "bronnen": bronnen,
        "vakken": vakken,
    }
    with open(os.path.join(OUT, "tk-syllabus-v6.json"), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=0)
    n = sum(len(v) for v in vakken.values())
    print(f"{len(vakken)} vakken, {n} regels, {sum(1 for v in vakken.values() for i in v if i['bron'])} met bron → sources/ecqb/tk-syllabus-v6.json")


if __name__ == "__main__":
    main(sys.argv[1])
